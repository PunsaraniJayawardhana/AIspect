"""
backend/scripts/validate_lexicon.py

Validates the G(s) ambiguity lexicon (backend/lexicons/ambiguous_terms.py)
against two independent human annotators' manual tagging of the same
held-out AC sample.

Expects an input spreadsheet with three columns:
    "AC text", "your flagged terms", "annotator2 flagged terms"
where the two annotation columns are comma-separated lists of terms
(case-insensitive), one row per AC, matching what the ADF parser's
explicit_ACs actually produced (not hand-retyped/cleaned-up text).

Computes:
  1. Human inter-rater agreement (Jaccard similarity, per-AC and averaged)
  2. Lexicon precision/recall/F1 against a STRICT ground truth (terms both
     annotators agreed on) and a LENIENT ground truth (terms either
     annotator flagged) -- reporting both gives a defensible range rather
     than a single, arguably-arbitrary choice of ground truth definition.
  3. A full false-positive / false-negative list per AC, for manually
     deciding which lexicon terms to keep, drop, or add.

Usage:
    python -m backend.scripts.validate_lexicon path/to/AC_ambiguity_annotation.xlsx
"""

import sys
from openpyxl import load_workbook
from backend.lexicons.ambiguous_terms import find_ambiguous_terms


def parse_terms(cell_value) -> set:
    """Turn a comma-separated cell into a lowercase, stripped set of terms."""
    if not cell_value:
        return set()
    return {t.strip().lower() for t in str(cell_value).split(",") if t.strip()}


def load_annotations(path: str):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = []
    for r in range(2, ws.max_row + 1):
        ac_text = ws.cell(row=r, column=1).value
        if not ac_text:
            continue
        a1 = parse_terms(ws.cell(row=r, column=2).value)
        a2 = parse_terms(ws.cell(row=r, column=3).value)
        rows.append((ac_text, a1, a2))
    return rows


def jaccard(a: set, b: set) -> float:
    """1.0 if both sets are empty (perfect agreement on 'nothing ambiguous')."""
    if not a and not b:
        return 1.0
    union = a | b
    if not union:
        return 1.0
    return len(a & b) / len(union)


def prf1(tp: int, fp: int, fn: int):
    precision = tp / (tp + fp) if (tp + fp) > 0 else 0.0
    recall = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    f1 = 2 * precision * recall / (precision + recall) if (precision + recall) > 0 else 0.0
    return precision, recall, f1


def main():
    if len(sys.argv) < 2:
        print("Usage: python -m backend.scripts.validate_lexicon path/to/annotations.xlsx")
        sys.exit(1)

    path = sys.argv[1]
    rows = load_annotations(path)
    print(f"Loaded {len(rows)} annotated ACs from {path}\n")

    # ── 1. Human inter-rater agreement ──
    jaccards = [jaccard(a1, a2) for _, a1, a2 in rows]
    mean_jaccard = sum(jaccards) / len(jaccards)
    print("=" * 70)
    print("HUMAN INTER-RATER AGREEMENT")
    print("=" * 70)
    print(f"Mean per-AC Jaccard similarity: {mean_jaccard:.4f}")
    print("(1.0 = perfect agreement on every AC; 0.0 = no overlap at all)\n")

    # ── 2. Lexicon vs. two ground-truth definitions ──
    strict_tp = strict_fp = strict_fn = 0
    lenient_tp = lenient_fp = lenient_fn = 0
    mismatches = []

    for ac_text, a1, a2 in rows:
        strict_truth = a1 & a2   # both agreed
        lenient_truth = a1 | a2  # either flagged

        lex_matches = {t.lower() for t in find_ambiguous_terms(ac_text)}

        s_tp = lex_matches & strict_truth
        s_fp = lex_matches - strict_truth
        s_fn = strict_truth - lex_matches
        strict_tp += len(s_tp); strict_fp += len(s_fp); strict_fn += len(s_fn)

        l_tp = lex_matches & lenient_truth
        l_fp = lex_matches - lenient_truth
        l_fn = lenient_truth - lex_matches
        lenient_tp += len(l_tp); lenient_fp += len(l_fp); lenient_fn += len(l_fn)

        if l_fp or l_fn:
            mismatches.append({
                "ac": ac_text,
                "lexicon_flagged": sorted(lex_matches),
                "false_positives": sorted(l_fp),   # lexicon flagged, no human did
                "false_negatives": sorted(l_fn),   # a human flagged, lexicon missed
            })

    print("=" * 70)
    print("LEXICON PRECISION / RECALL / F1")
    print("=" * 70)
    for label, tp, fp, fn in [
        ("STRICT ground truth (both annotators agreed)", strict_tp, strict_fp, strict_fn),
        ("LENIENT ground truth (either annotator flagged)", lenient_tp, lenient_fp, lenient_fn),
    ]:
        p, r, f1 = prf1(tp, fp, fn)
        print(f"\n{label}:")
        print(f"  TP={tp}  FP={fp}  FN={fn}")
        print(f"  Precision={p:.4f}  Recall={r:.4f}  F1={f1:.4f}")

    print()
    print("=" * 70)
    print(f"MISMATCHES FOR REVIEW ({len(mismatches)} ACs with at least one disagreement)")
    print("=" * 70)
    for m in mismatches:
        print(f"\nAC: {m['ac'][:100]}")
        if m["false_positives"]:
            print(f"  FALSE POSITIVES (lexicon flagged, no human agreed): {m['false_positives']}")
        if m["false_negatives"]:
            print(f"  FALSE NEGATIVES (a human flagged, lexicon missed):  {m['false_negatives']}")


if __name__ == "__main__":
    main()