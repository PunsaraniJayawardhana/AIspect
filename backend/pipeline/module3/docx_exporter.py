import os

from docx import Document


def _normalized_rows(tests):
    rows = []
    for test in tests or []:
        rows.append(
            {
                "tc_id": test.get("tc_id", ""),
                "mode": test.get("mode", ""),
                "scenario": test.get("scenario", ""),
                "priority": test.get("priority", ""),
                "tpri": test.get("tpri_score") if test.get("tpri_score") is not None else "",
                "rank": test.get("priority_rank") if test.get("priority_rank") is not None else "",
                "steps": "; ".join(test.get("steps", [])),
                "expected_result": test.get("expected_result", ""),
            }
        )
    return rows


def _ensure_parent(path):
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)


def export_to_docx(tests, story_key, output_path=None):
    output_path = output_path or f"backend/scripts/fixtures/{story_key}_test_cases.docx"
    _ensure_parent(output_path)

    document = Document()
    document.add_heading(f"Test Cases — {story_key}", level=1)

    table = document.add_table(rows=1, cols=8)
    headers = [
        "TC ID",
        "Mode",
        "Scenario",
        "Priority",
        "TPRI",
        "Rank",
        "Steps",
        "Expected Result",
    ]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header

    for row in _normalized_rows(tests):
        cells = table.add_row().cells
        cells[0].text = str(row["tc_id"])
        cells[1].text = str(row["mode"])
        cells[2].text = str(row["scenario"])
        cells[3].text = str(row["priority"])
        cells[4].text = str(row["tpri"])
        cells[5].text = str(row["rank"])
        cells[6].text = str(row["steps"])
        cells[7].text = str(row["expected_result"])

    document.save(output_path)
    return output_path


def export_to_markdown(tests, story_key, output_path=None):
    output_path = output_path or f"backend/scripts/fixtures/{story_key}_test_cases.md"
    _ensure_parent(output_path)

    rows = _normalized_rows(tests)
    lines = [
        f"# Test Cases — {story_key}",
        "",
        "| TC ID | Mode | Scenario | Priority | TPRI | Rank | Steps | Expected Result |",
        "| --- | --- | --- | --- | --- | --- | --- | --- |",
    ]

    for row in rows:
        lines.append(
            "| {tc_id} | {mode} | {scenario} | {priority} | {tpri} | {rank} | {steps} | {expected_result} |".format(
                tc_id=row["tc_id"],
                mode=row["mode"],
                scenario=row["scenario"].replace("|", "\\|")[:120],
                priority=row["priority"],
                tpri=row["tpri"],
                rank=row["rank"],
                steps=row["steps"].replace("|", "\\|")[:120],
                expected_result=row["expected_result"].replace("|", "\\|")[:120],
            )
        )

    with open(output_path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines) + "\n")

    return output_path
